import os
import io
import calendar
from datetime import datetime, date
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
from typing import Optional, List, Any
import psycopg2
from psycopg2.extras import RealDictCursor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/py/assets-tracker", tags=["Assets Own EQ Parts Tracker"])

def get_db_connection():
    return psycopg2.connect(
        dbname=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASS"),
        host=os.getenv("DB_HOST", "127.0.0.1"),
        port=os.getenv("DB_PORT", "5432")
    )

class NewAssetTypePayload(BaseModel):
    type_name: str

class NewAssetMasterPayload(BaseModel):
    asset_type_id: int
    asset_code: str
    purchase_date: Optional[str] = None
    purchase_cost: Optional[float] = 0.0
    remaining_purchase_price: Optional[float] = 0.0

class EditAssetMasterCostPayload(BaseModel):
    asset_id: int
    purchase_date: Optional[str] = None
    purchase_cost: Optional[float] = 0.0
    remaining_purchase_price: Optional[float] = 0.0

class MonthRecordItem(BaseModel):
    month: int
    site_name: Optional[str] = ""
    plate_no: Optional[str] = ""
    off_site_exp: Optional[float] = 0.0
    on_site_exp: Optional[float] = 0.0
    revenue: Optional[float] = 0.0

class SaveMonthRecordsPayload(BaseModel):
    asset_id: int
    year: int
    records: List[MonthRecordItem]

@router.get("/types")
def get_asset_types():
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute("SELECT id, type_name FROM assets_types ORDER BY id ASC;")
        rows = cur.fetchall()
        cur.close()
        conn.close()
        return {"success": True, "types": rows}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/types/add")
def add_asset_type(payload: NewAssetTypePayload):
    try:
        name = payload.type_name.strip()
        if not name:
            raise HTTPException(status_code=400, detail="Type name cannot be empty")
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute(
            "INSERT INTO assets_types (type_name) VALUES (%s) ON CONFLICT (type_name) DO NOTHING RETURNING id;",
            (name,)
        )
        res = cur.fetchone()
        conn.commit()
        cur.close()
        conn.close()
        return {"success": True, "message": "Asset type added successfully", "id": res[0] if res else None}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/master/add")
def add_asset_master(payload: NewAssetMasterPayload):
    try:
        conn = get_db_connection()
        cur = conn.cursor()
        p_date = payload.purchase_date if payload.purchase_date else None
        cur.execute(
            """
            INSERT INTO assets_master (asset_type_id, asset_code, purchase_date, purchase_cost, remaining_purchase_price)
            VALUES (%s, %s, %s, %s, %s)
            RETURNING id;
            """,
            (payload.asset_type_id, payload.asset_code.strip(), p_date, payload.purchase_cost, payload.remaining_purchase_price)
        )
        new_id = cur.fetchone()[0]
        conn.commit()
        cur.close()
        conn.close()
        return {"success": True, "id": new_id, "message": "Asset created successfully"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/data")
def get_assets_data():
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        
        cur.execute("""
            SELECT 
                m.id, 
                m.asset_type_id,
                t.type_name,
                m.asset_code, 
                m.purchase_date, 
                COALESCE(m.purchase_cost, 0) AS purchase_cost,
                COALESCE(m.remaining_purchase_price, 0) AS remaining_purchase_price
            FROM assets_master m
            LEFT JOIN assets_types t ON m.asset_type_id = t.id
            ORDER BY t.type_name ASC NULLS LAST, m.purchase_date ASC NULLS LAST, m.id ASC;
        """)
        assets = cur.fetchall()

        cur.execute("""
            SELECT 
                asset_id, year, month, 
                COALESCE(site_name, '') as site_name, 
                COALESCE(plate_no, '') as plate_no, 
                COALESCE(off_site_exp, 0) as off_site_exp, 
                COALESCE(on_site_exp, 0) as on_site_exp, 
                COALESCE(revenue, 0) as revenue
            FROM asset_monthly_records;
        """)
        records = cur.fetchall()
        cur.close()
        conn.close()

        rec_map = {}
        all_months = set()
        for r in records:
            k = f"{r['asset_id']}_{r['year']}_{r['month']}"
            rec_map[k] = r
            if r['site_name'] or r['plate_no'] or r['off_site_exp'] > 0 or r['on_site_exp'] > 0 or r['revenue'] > 0:
                all_months.add((r['year'], r['month']))

        today = date.today()
        valid_dates = [a["purchase_date"] for a in assets if a["purchase_date"]]
        if valid_dates:
            earliest = min(valid_dates)
            min_y, min_m = earliest.year, earliest.month
        elif all_months:
            min_y, min_m = min(all_months)
        else:
            min_y, min_m = today.year, 1

        timeline = []
        cur_y, cur_m = min_y, min_m
        while (cur_y < today.year) or (cur_y == today.year and cur_m <= today.month):
            timeline.append({
                "year": cur_y,
                "month": cur_m,
                "label": f"{calendar.month_abbr[cur_m]} {cur_y}"
            })
            if cur_m == 12:
                cur_y += 1
                cur_m = 1
            else:
                cur_m += 1

        formatted_assets = []
        for a in assets:
            a_id = a["id"]
            m_data = []
            for tm in timeline:
                k = f"{a_id}_{tm['year']}_{tm['month']}"
                r_item = rec_map.get(k, {
                    "site_name": "",
                    "plate_no": "",
                    "off_site_exp": 0,
                    "on_site_exp": 0,
                    "revenue": 0
                })
                m_data.append({
                    "site_name": r_item["site_name"],
                    "plate_no": r_item["plate_no"],
                    "off_site_exp": float(r_item["off_site_exp"]),
                    "on_site_exp": float(r_item["on_site_exp"]),
                    "revenue": float(r_item["revenue"])
                })
            formatted_assets.append({
                **a,
                "months_data": m_data
            })

        return {
            "success": True,
            "assets": formatted_assets,
            "timeline": timeline
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/master/edit-costs")
def edit_asset_master_costs(payload: EditAssetMasterCostPayload):
    try:
        conn = get_db_connection()
        cur = conn.cursor()
        p_date = payload.purchase_date if payload.purchase_date else None
        cur.execute(
            """
            UPDATE assets_master
            SET purchase_date = %s,
                purchase_cost = %s,
                remaining_purchase_price = %s,
                updated_at = CURRENT_TIMESTAMP
            WHERE id = %s;
            """,
            (p_date, payload.purchase_cost, payload.remaining_purchase_price, payload.asset_id)
        )
        conn.commit()
        cur.close()
        conn.close()
        return {"success": True, "message": "Asset purchase details updated successfully!"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/get-year-records")
def get_year_records(asset_id: int, year: int):
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute("""
            SELECT month, site_name, plate_no, off_site_exp, on_site_exp, revenue
            FROM asset_monthly_records
            WHERE asset_id = %s AND year = %s
            ORDER BY month ASC;
        """, (asset_id, year))
        rows = cur.fetchall()
        cur.close()
        conn.close()

        rec_dict = {r["month"]: r for r in rows}
        result = []
        for m in range(1, 13):
            if m in rec_dict:
                result.append(rec_dict[m])
            else:
                result.append({
                    "month": m,
                    "site_name": "",
                    "plate_no": "",
                    "off_site_exp": 0.0,
                    "on_site_exp": 0.0,
                    "revenue": 0.0
                })
        return {"success": True, "records": result}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/save-year-records")
def save_year_records(payload: SaveMonthRecordsPayload):
    try:
        conn = get_db_connection()
        cur = conn.cursor()
        for r in payload.records:
            cur.execute("""
                INSERT INTO asset_monthly_records (
                    asset_id, year, month, site_name, plate_no, off_site_exp, on_site_exp, revenue, updated_at
                ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
                ON CONFLICT (asset_id, year, month)
                DO UPDATE SET
                    site_name = EXCLUDED.site_name,
                    plate_no = EXCLUDED.plate_no,
                    off_site_exp = EXCLUDED.off_site_exp,
                    on_site_exp = EXCLUDED.on_site_exp,
                    revenue = EXCLUDED.revenue,
                    updated_at = CURRENT_TIMESTAMP;
            """, (
                payload.asset_id, payload.year, r.month,
                r.site_name.strip() if r.site_name else "",
                r.plate_no.strip() if r.plate_no else "",
                float(r.off_site_exp or 0),
                float(r.on_site_exp or 0),
                float(r.revenue or 0)
            ))
        conn.commit()
        cur.close()
        conn.close()
        return {"success": True, "message": "Monthly records saved successfully"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/export-excel")
def export_assets_excel(
    mode: str = "single",
    start_year: Optional[int] = None,
    start_month: Optional[int] = None,
    end_year: Optional[int] = None,
    end_month: Optional[int] = None
):
    try:
        data_res = get_assets_data()
        assets = data_res["assets"]
        full_timeline = data_res["timeline"]

        wb = openpyxl.Workbook()
        default_sheet = wb.active

        head_fill_main = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        head_fill_month = PatternFill(start_color="C6D9F1", end_color="C6D9F1", fill_type="solid")
        head_fill_sub = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
        rev_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        head_font = Font(name="Segoe UI", size=11, bold=True, color="000000")
        regular_font = Font(name="Segoe UI", size=10)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='C0C0C0'),
            right=Side(style='thin', color='C0C0C0'),
            top=Side(style='thin', color='C0C0C0'),
            bottom=Side(style='thin', color='C0C0C0')
        )

        base_headers = [
            "SN", "Asset Type", "Asset Code", "Date of Purchase", "Cost of Purchase", "Remaining Purchase Price"
        ]

        def build_sheet(ws, sheet_timeline):
            ws.row_dimensions[1].height = 25
            ws.row_dimensions[2].height = 25

            for c_idx, h_text in enumerate(base_headers, 1):
                ws.merge_cells(start_row=1, start_column=c_idx, end_row=2, end_column=c_idx)
                c = ws.cell(row=1, column=c_idx, value=h_text)
                c.fill = head_fill_main
                c.font = head_font
                c.alignment = center_align

            col_cursor = len(base_headers) + 1
            for tm in sheet_timeline:
                ws.merge_cells(start_row=1, start_column=col_cursor, end_row=1, end_column=col_cursor + 4)
                c_top = ws.cell(row=1, column=col_cursor, value=tm["label"])
                c_top.fill = head_fill_month
                c_top.font = head_font
                c_top.alignment = center_align

                sub_headers = ["Site", "Plate No", "Off Site Exp", "On site", "Revenue"]
                for sub_idx, sub_label in enumerate(sub_headers):
                    c_sub = ws.cell(row=2, column=col_cursor + sub_idx, value=sub_label)
                    c_sub.fill = head_fill_sub
                    c_sub.font = head_font
                    c_sub.alignment = center_align
                col_cursor += 5

            total_cols = col_cursor - 1
            for r in [1, 2]:
                for c in range(1, total_cols + 1):
                    ws.cell(row=r, column=c).border = thin_border

            curr_row = 3
            for idx, item in enumerate(assets, 1):
                ws.row_dimensions[curr_row].height = 24
                p_date = item.get("purchase_date")
                p_date_obj = None
                if p_date:
                    try:
                        p_date_obj = datetime.strptime(str(p_date)[:10], "%Y-%m-%d").date()
                    except Exception:
                        p_date_obj = None

                ws.cell(row=curr_row, column=1, value=idx).alignment = center_align
                ws.cell(row=curr_row, column=2, value=item.get("type_name") or "").alignment = center_align
                ws.cell(row=curr_row, column=3, value=item.get("asset_code") or "").alignment = center_align

                c_date = ws.cell(row=curr_row, column=4, value=p_date_obj)
                c_date.alignment = center_align
                if p_date_obj:
                    c_date.number_format = 'dd-mmm-yy'

                c_cost = ws.cell(row=curr_row, column=5, value=float(item.get("purchase_cost") or 0))
                c_cost.alignment = center_align
                c_cost.number_format = '#,##0.00'

                c_rem = ws.cell(row=curr_row, column=6, value=float(item.get("remaining_purchase_price") or 0))
                c_rem.alignment = center_align
                c_rem.number_format = '#,##0.00'

                m_data_list = item.get("months_data", [])
                for t_idx, tm in enumerate(sheet_timeline):
                    orig_t_idx = next((i for i, x in enumerate(full_timeline) if x["year"] == tm["year"] and x["month"] == tm["month"]), -1)
                    m_info = m_data_list[orig_t_idx] if orig_t_idx != -1 and orig_t_idx < len(m_data_list) else {
                        "site_name": "", "plate_no": "", "off_site_exp": 0, "on_site_exp": 0, "revenue": 0
                    }

                    s_col = 7 + (t_idx * 5)
                    ws.cell(row=curr_row, column=s_col, value=m_info.get("site_name") or "").alignment = left_align
                    ws.cell(row=curr_row, column=s_col + 1, value=m_info.get("plate_no") or "").alignment = center_align

                    off_val = float(m_info.get("off_site_exp") or 0)
                    c_off = ws.cell(row=curr_row, column=s_col + 2, value=off_val if off_val > 0 else "")
                    c_off.alignment = center_align
                    if off_val > 0: c_off.number_format = '#,##0.00'

                    on_val = float(m_info.get("on_site_exp") or 0)
                    c_on = ws.cell(row=curr_row, column=s_col + 3, value=on_val if on_val > 0 else "")
                    c_on.alignment = center_align
                    if on_val > 0: c_on.number_format = '#,##0.00'

                    rev_val = float(m_info.get("revenue") or 0)
                    c_rev = ws.cell(row=curr_row, column=s_col + 4, value=rev_val if rev_val > 0 else "")
                    c_rev.alignment = center_align
                    c_rev.fill = rev_fill
                    if rev_val > 0: c_rev.number_format = '#,##0.00'

                for c_num in range(1, total_cols + 1):
                    cell = ws.cell(row=curr_row, column=c_num)
                    cell.border = thin_border
                    if cell.font == Font(): cell.font = regular_font
                curr_row += 1

            last_col_letter = get_column_letter(total_cols)
            ws.auto_filter.ref = f"A2:{last_col_letter}{curr_row - 1}"
            ws.sheet_view.zoomScale = 80
            ws.freeze_panes = 'A3'

            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    val = cell.value
                    if val:
                        val_str = val.strftime('%d-%b-%y') if isinstance(val, (datetime, date)) else str(val)
                        max_len = max(max_len, len(val_str))
                ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        if mode == "yearly":
            years_list = sorted(list(set(tm["year"] for tm in full_timeline)))
            if not years_list: years_list = [date.today().year]
            for y_idx, yr in enumerate(years_list):
                yr_timeline = [tm for tm in full_timeline if tm["year"] == yr]
                ws = default_sheet if y_idx == 0 else wb.create_sheet(title=str(yr))
                ws.title = str(yr)
                build_sheet(ws, yr_timeline)
        elif mode == "period":
            s_y = start_year or full_timeline[0]["year"]
            s_m = start_month or full_timeline[0]["month"]
            e_y = end_year or full_timeline[-1]["year"]
            e_m = end_month or full_timeline[-1]["month"]

            period_timeline = []
            for tm in full_timeline:
                t_val = tm["year"] * 100 + tm["month"]
                if (s_y * 100 + s_m) <= t_val <= (e_y * 100 + e_m):
                    period_timeline.append(tm)
            if not period_timeline: period_timeline = full_timeline
            default_sheet.title = "Period Export"
            build_sheet(default_sheet, period_timeline)
        else:
            default_sheet.title = "Assets Tracker"
            build_sheet(default_sheet, full_timeline)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"Assets_Tracker_{date.today().strftime('%Y_%m_%d')}.xlsx"
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))