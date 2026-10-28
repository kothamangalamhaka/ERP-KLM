import os
import io
from datetime import datetime
import psycopg2
from psycopg2.extras import RealDictCursor
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/py/johns-report", tags=["Johns Report"])

def get_db_connection():
    return psycopg2.connect(
        dbname=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASS"),
        host=os.getenv("DB_HOST", "localhost"),
        port=os.getenv("DB_PORT", "5432")
    )

class MonthItem(BaseModel):
    m: str
    y: int

class FetchBatchPayload(BaseModel):
    months: list[MonthItem]

class ExcelSheetPayload(BaseModel):
    sheet_name: str
    rows: list[list]
    custom_headers: dict = {}

class ExportPayload(BaseModel):
    file_name: str
    sheets: list[ExcelSheetPayload]

class VerifyCodePayload(BaseModel):
    code: str

class CreateBackupPayload(BaseModel):
    backup_name: str
    period_range: str
    report_data: dict

class DeleteBackupPayload(BaseModel):
    pass_key: str

@router.post("/verify-code")
def verify_security_code(payload: VerifyCodePayload):
    valid_code = os.getenv("JOHN_REPORT_CODE", "John_Report")
    if payload.code == valid_code:
        return {"success": True}
    return {"success": False, "message": "Invalid Security Code!"}

@router.post("/create-backup")
def create_report_backup(payload: CreateBackupPayload):
    try:
        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute(
            """
            INSERT INTO johns_report_backups (backup_name, period_range, report_data)
            VALUES (%s, %s, %s::jsonb)
            RETURNING id, created_at
            """,
            (payload.backup_name, payload.period_range, psycopg2.extras.Json(payload.report_data))
        )
        res = cur.fetchone()
        conn.commit()
        cur.close()
        conn.close()
        return {"success": True, "message": "Backup created successfully!", "id": res[0], "created_at": res[1]}
    except Exception as e:
        return {"success": False, "message": str(e)}

@router.get("/list-backups")
def get_backup_list():
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            """
            SELECT id, backup_name, period_range, TO_CHAR(created_at, 'YYYY-MM-DD HH24:MI:SS') as created_at
            FROM johns_report_backups
            ORDER BY id DESC
            """
        )
        backups = cur.fetchall()
        cur.close()
        conn.close()
        return {"success": True, "data": backups}
    except Exception as e:
        return {"success": False, "message": str(e)}

@router.get("/get-backup/{backup_id}")
def get_backup_detail(backup_id: int):
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        cur.execute(
            "SELECT id, backup_name, period_range, report_data, TO_CHAR(created_at, 'YYYY-MM-DD HH24:MI:SS') as created_at FROM johns_report_backups WHERE id = %s",
            (backup_id,)
        )
        backup = cur.fetchone()
        cur.close()
        conn.close()
        if not backup:
            return {"success": False, "message": "Backup not found."}
        return {"success": True, "data": backup}
    except Exception as e:
        return {"success": False, "message": str(e)}

@router.post("/delete-backup/{backup_id}")
def delete_report_backup(backup_id: int, payload: DeleteBackupPayload):
    try:
        valid_key = os.getenv("JOHN_REPORT_DELETE_KEY", "John_Delete")
        if payload.pass_key != valid_key:
            return {"success": False, "message": "Invalid Delete Pass Key!"}

        conn = get_db_connection()
        cur = conn.cursor()
        cur.execute("DELETE FROM johns_report_backups WHERE id = %s", (backup_id,))
        conn.commit()
        cur.close()
        conn.close()
        return {"success": True, "message": "Backup deleted successfully!"}
    except Exception as e:
        return {"success": False, "message": str(e)}

@router.post("/fetch-data")
def fetch_report_data(payload: FetchBatchPayload):
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)

        full_months = [f"{item.m} {item.y}" for item in payload.months]

        cur.execute("SELECT plate_no, owner_name, site_name, vehicle_type, vat FROM timesheet_vehicles")
        vehicles = cur.fetchall()

        cur.execute("SELECT plate_no, site_name, work_start_date, work_end_date, rate, field_co, site_co, vehicle_type, status FROM vehicle_site_log")
        sites = cur.fetchall()

        cur.execute("SELECT plate_no, driver_name, work_start_date, work_end_date FROM vehicle_driver_log")
        drivers = cur.fetchall()

        cur.execute("SELECT plate_no, record_date, calc_time, bd, month, year FROM timesheet_daily_records WHERE (month, year::int) IN %s", (tuple((item.m, item.y) for item in payload.months),))
        timesheets = cur.fetchall()

        cur.execute("SELECT * FROM invoice_records WHERE month = ANY(%s)", (full_months,))
        invoices = cur.fetchall()

        cur.execute("SELECT * FROM billing_records WHERE billing_month = ANY(%s)", (full_months,))
        billing = cur.fetchall()

        cur.execute("SELECT plate_no, owner_name, vat, work_start_date, work_end_date, status FROM vehicle_owner_log")
        owners = cur.fetchall()

        cur.execute("SELECT plate_no, site_name, rate, work_start_date, work_end_date, status FROM vehicle_rate_log")
        rates = cur.fetchall()

        cur.execute("SELECT old_plate_no, new_plate_no, TO_CHAR(change_date, 'YYYY-MM-DD') as change_date FROM vehicle_plate_log ORDER BY change_date ASC")
        plate_changes = cur.fetchall()

        cur.close()
        conn.close()

        return {
            "success": True,
            "vehicles": vehicles,
            "sites": sites,
            "drivers": drivers,
            "timesheets": timesheets,
            "invoices": invoices,
            "billing": billing,
            "owners": owners,
            "rates": rates,
            "plateChanges": plate_changes
        }
    except Exception as e:
        return {"success": False, "message": str(e)}

@router.post("/export-excel")
def export_johns_excel(payload: ExportPayload):
    try:
        output = io.BytesIO()
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )
        hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell_font = Font(name="Calibri", size=10)
        
        dark_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        calc_fill = PatternFill(start_color="0F766E", end_color="0F766E", fill_type="solid")

        diff_fill = PatternFill(start_color="BE123C", end_color="BE123C", fill_type="solid")
        pos_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        neg_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

        amt_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
        inv_fill = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")
        pl_fill = PatternFill(start_color="4338CA", end_color="4338CA", fill_type="solid")

        top_headers = [
            ("VEHICLE INFO", 1, 10, dark_fill),
            ("Logsheet Collected", 11, 12, dark_fill),
            ("Bill Summary", 13, 14, dark_fill),
            ("Supplier", 15, 16, dark_fill),
            ("For Driver", 17, 18, dark_fill),
            ("CALCULATIONS", 19, 21, calc_fill),
            ("DIFFERENCE", 22, 25, diff_fill),
            ("SUPPLIER PAYMENTS", 26, 31, amt_fill),
            ("INVOICE AMOUNT", 32, 38, inv_fill),
            ("P&L", 39, 39, pl_fill)
        ]

        sub_headers = [
            "SN", "Site", "Rate", "VAT", "Owner Name", "Driver Name",
            "Site Co", "Field Co", "Vehicle Type", "Plate No",
            "NR", "OT", "NR", "OT", "NR", "OT", "OT", "AMT",
            "Logsheet CF", "Bill Summary CF", "Supplier CF",
            "Log - Supplier", "Bill - Supplier", "Bill - Log", "Notes",
            "NR Rate", "OT Rate", "Rent", "VAT", "Total", "Net Total",
            "Subtotal", "Deductions", "Total", "Reason", "CRN No", "Invoice No", "Bill No",
            "Profit / Loss"
        ]

        for s_idx, sheet_data in enumerate(payload.sheets):
            ws = wb.create_sheet(title=sheet_data.sheet_name[:31])
            ws.views.sheetView[0].showGridLines = True

            for title, sc, ec, fill in top_headers:
                ws.merge_cells(start_row=1, start_column=sc, end_row=1, end_column=ec)
                for col in range(sc, ec + 1):
                    cell = ws.cell(row=1, column=col)
                    cell.fill = fill
                    cell.border = thin_border
                    if col == sc:
                        cell.value = title
                        cell.font = hdr_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")

            for col_idx, col_name in enumerate(sub_headers, 1):
                cell = ws.cell(row=2, column=col_idx, value=col_name)
                if col_idx in [19, 20, 21]:
                    cell.fill = calc_fill
                elif 22 <= col_idx <= 25:
                    cell.fill = diff_fill
                elif 26 <= col_idx <= 31:
                    cell.fill = amt_fill
                elif 32 <= col_idx <= 38:
                    cell.fill = inv_fill
                elif col_idx == 39:
                    cell.fill = pl_fill
                else:
                    cell.fill = dark_fill
                cell.font = hdr_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = thin_border

            for col_idx, col_name in enumerate(sub_headers, 1):
                cell = ws.cell(row=2, column=col_idx)
                if col_idx == 22 and "diff_log_sup" in sheet_data.custom_headers:
                    cell.value = sheet_data.custom_headers["diff_log_sup"]
                    cell.font = Font(name="Calibri", size=8, bold=False, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                elif col_idx == 23 and "diff_bill_sup" in sheet_data.custom_headers:
                    cell.value = sheet_data.custom_headers["diff_bill_sup"]
                    cell.font = Font(name="Calibri", size=8, bold=False, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                elif col_idx == 24 and "diff_bill_log" in sheet_data.custom_headers:
                    cell.value = sheet_data.custom_headers["diff_bill_log"]
                    cell.font = Font(name="Calibri", size=8, bold=False, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            ws.row_dimensions[1].height = 24
            ws.row_dimensions[2].height = 36

            for r_idx, row_values in enumerate(sheet_data.rows, 3):
                for c_idx, val in enumerate(row_values, 1):
                    cell_val = val
                    if c_idx in [22, 23, 24]:
                        try:
                            if val not in [None, "", "-"] and float(val) == 0:
                                cell_val = ""
                        except (ValueError, TypeError):
                            pass

                    cell = ws.cell(row=r_idx, column=c_idx, value=cell_val)
                    cell.font = cell_font
                    cell.border = thin_border
                    if c_idx in [22, 23, 24] and cell_val not in [None, "", "-"]:
                        try:
                            num_val = float(cell_val)
                            if num_val > 0:
                                cell.fill = pos_fill
                            elif num_val < 0:
                                cell.fill = neg_fill
                        except (ValueError, TypeError):
                            pass
                    if c_idx in [22, 23, 24] and "\n" in str(cell_val or ""):
                        cell.font = Font(name="Calibri", size=9, bold=True)
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    elif c_idx in [2, 5, 6, 7, 8, 9]:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

            for col in ws.columns:
                c_num = col[0].column
                if c_num in [22, 23, 24]:
                    col_letter = get_column_letter(c_num)
                    ws.column_dimensions[col_letter].width = 17
                else:
                    max_len = max(len(str(cell.value or '')) for cell in col)
                    col_letter = get_column_letter(c_num)
                    ws.column_dimensions[col_letter].width = max(max_len + 3, 9)

            last_row = max(len(sheet_data.rows) + 2, 2)
            last_col_letter = get_column_letter(len(sub_headers))
            ws.auto_filter.ref = f"A2:{last_col_letter}{last_row}"
            ws.sheet_view.zoomScale = 80

        wb.save(output)
        output.seek(0)
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={payload.file_name}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))