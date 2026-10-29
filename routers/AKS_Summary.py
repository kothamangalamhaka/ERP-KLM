import os
import calendar
from datetime import datetime, date
from fastapi import APIRouter, HTTPException
import io
from fastapi.responses import Response
from pydantic import BaseModel
from typing import Optional, List, Any
import psycopg2
from psycopg2.extras import RealDictCursor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter(prefix="/py/aks-summary", tags=["AKS Summary Tracker"])

def get_db_connection():
    return psycopg2.connect(
        dbname=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASS"),
        host=os.getenv("DB_HOST", "127.0.0.1"),
        port=os.getenv("DB_PORT", "5432")
    )

def init_aks_summary_columns():
    conn = get_db_connection()
    cur = conn.cursor()
    cur.execute("""
        ALTER TABLE equipment_aks_details ADD COLUMN IF NOT EXISTS asset_type VARCHAR(100) DEFAULT '';
        ALTER TABLE equipment_aks_details ADD COLUMN IF NOT EXISTS model VARCHAR(100) DEFAULT '';
        ALTER TABLE equipment_aks_details ADD COLUMN IF NOT EXISTS make VARCHAR(100) DEFAULT '';
    """)
    conn.commit()
    cur.close()
    conn.close()

try:
    init_aks_summary_columns()
except Exception as e:
    print(f"AKS Summary Table init note: {e}")

class AKSSummaryCellUpdate(BaseModel):
    equipment_id: int
    field: str
    value: Any

class AKSSummarySavePayload(BaseModel):
    updates: List[AKSSummaryCellUpdate]

ALLOWED_FIELDS = {
    "asset_type",
    "model",
    "make",
    "owner_name"
}

def clean_site_name(raw_site: str) -> str:
    if not raw_site:
        return ""
    import re
    clean = str(raw_site).strip()
    unwanted = [
        r'\bL\s*&\s*T\s+Solar\b', r'\bL\s*&\s*T\b', r'\bHydro\b',
        r'\(?\s*re\s*running\s*\)?', r'\brerunning\b', r'\b22\b', r'\bZ\b', r'\bHaka\b',
        r'\bAl\s*Joda\b', r'\bAljoda\b', r'\bMasar\s*Wheels\b',
        r'\bMasar\b', r'\bWe\s*1\b', r'\bWe1\b', r'\bWe-1\b'
    ]
    for p in unwanted:
        clean = re.sub(p, ' ', clean, flags=re.IGNORECASE)
    clean = re.sub(r'[-–—_/\\,|()]+', ' ', clean)
    return re.sub(r'\s+', ' ', clean).strip()

@router.get("/data")
def get_aks_summary_data():
    try:
        conn = get_db_connection()
        cur = conn.cursor(cursor_factory=RealDictCursor)
        
        cur.execute("""
            SELECT 
                e.id,
                e.plate_no,
                e.purchase_date,
                COALESCE(e.purchase_cost, 0) AS purchase_cost,
                COALESCE(a.off_site_insurance, 0) AS off_site_insurance,
                COALESCE(a.on_site_insurance, 0) AS on_site_insurance,
                COALESCE(a.asset_type, '') AS asset_type,
                COALESCE(a.model, '') AS model,
                COALESCE(a.make, '') AS make,
                COALESCE(a.owner_name, '') AS owner_name
            FROM equipments e
            LEFT JOIN equipment_aks_details a ON e.id = a.equipment_id
            ORDER BY e.purchase_date ASC NULLS LAST, e.id ASC;
        """)
        equipments = cur.fetchall()

        cur.execute("""
            SELECT plate_no, site_name, work_start_date, work_end_date 
            FROM vehicle_site_log 
            ORDER BY work_start_date ASC NULLS LAST, id ASC;
        """)
        site_logs = cur.fetchall()

        cur.execute("""
            SELECT 
                equipment_id, year, month, 
                COALESCE(maintenance_cost, 0) AS maintenance_cost
            FROM equipment_monthly_logs;
        """)
        maint_logs = cur.fetchall()

        cur.close()
        conn.close()

        maint_map = {}
        all_logged_months = set()
        for ml in maint_logs:
            k = f"{ml['equipment_id']}_{ml['year']}_{ml['month']}"
            m_cost = float(ml["maintenance_cost"])
            maint_map[k] = m_cost
            if m_cost > 0:
                all_logged_months.add((ml["year"], ml["month"]))

        for s in site_logs:
            if s["work_start_date"]:
                all_logged_months.add((s["work_start_date"].year, s["work_start_date"].month))

        today = date.today()
        valid_p_dates = [eq["purchase_date"] for eq in equipments if eq["purchase_date"]]
        if valid_p_dates:
            earliest_p = min(valid_p_dates)
            min_y, min_m = earliest_p.year, earliest_p.month
        elif all_logged_months:
            min_y, min_m = min(all_logged_months)
        else:
            min_y, min_m = today.year, today.month

        timeline = []
        cur_y, cur_m = min_y, min_m
        while (cur_y < today.year) or (cur_y == today.year and cur_m <= today.month):
            m_abbr = calendar.month_abbr[cur_m]
            timeline.append({
                "year": cur_y,
                "month": cur_m,
                "label": f"{m_abbr} {cur_y}"
            })
            if cur_m == 12:
                cur_y += 1
                cur_m = 1
            else:
                cur_m += 1

        records = []
        for eq in equipments:
            eq_id = eq["id"]
            p_clean = (eq["plate_no"] or "").strip().upper()
            matched_sites = [s for s in site_logs if (s["plate_no"] or "").strip().upper() == p_clean]

            # First site assignment
            first_site_date = None
            first_site_name = ""
            if matched_sites:
                sorted_sites = sorted([s for s in matched_sites if s["work_start_date"]], key=lambda x: x["work_start_date"])
                if sorted_sites:
                    first_site_date = sorted_sites[0]["work_start_date"]
                    first_site_name = clean_site_name(sorted_sites[0]["site_name"]) or sorted_sites[0]["site_name"]

            # Calculate Off-Site MTC
            site_intervals = []
            for s in matched_sites:
                s_start = s["work_start_date"]
                s_end = s["work_end_date"]
                if s_start:
                    site_intervals.append({
                        "start": s_start,
                        "end": s_end
                    })

            calculated_off_site_mtc = 0.0
            on_site_maint_map = {}
            all_eq_years_months = set((ml["year"], ml["month"]) for ml in maint_logs if ml["equipment_id"] == eq_id)

            for yr, mn in all_eq_years_months:
                total_mtc = maint_map.get(f"{eq_id}_{yr}_{mn}", 0.0)
                if total_mtc <= 0:
                    continue

                _, num_days = calendar.monthrange(yr, mn)
                daily_cost = total_mtc / num_days

                for day in range(1, num_days + 1):
                    current_d = date(yr, mn, day)
                    is_on_site = False
                    for interval in site_intervals:
                        st = interval["start"]
                        en = interval["end"]
                        if en:
                            if st <= current_d <= en:
                                is_on_site = True
                                break
                        else:
                            if current_d >= st:
                                is_on_site = True
                                break

                    if not is_on_site:
                        calculated_off_site_mtc += daily_cost
                    else:
                        on_site_maint_map[f"{yr}_{mn}"] = on_site_maint_map.get(f"{yr}_{mn}", 0.0) + daily_cost

            total_insurance = float(eq.get("off_site_insurance") or 0) + float(eq.get("on_site_insurance") or 0)
            net_off_site_mtc = max(0.0, calculated_off_site_mtc - float(eq.get("off_site_insurance") or 0))

            # Monthly MTC List (Single column per month)
            month_mtc_data = []
            for tm in timeline:
                yr, mn = tm["year"], tm["month"]
                mtc_val = on_site_maint_map.get(f"{yr}_{mn}", 0.0)
                month_mtc_data.append(round(mtc_val, 2))

            records.append({
                "id": eq["id"],
                "asset_type": eq["asset_type"],
                "model": eq["model"],
                "make": eq["make"],
                "owner_name": eq["owner_name"],
                "plate_no": eq["plate_no"],
                "purchase_date": eq["purchase_date"],
                "purchase_cost": float(eq["purchase_cost"]),
                "mtc_during_purchase": round(net_off_site_mtc, 2),
                "total_insurance": round(total_insurance, 2),
                "reached_site_date": first_site_date,
                "first_site": first_site_name,
                "monthly_mtc": month_mtc_data
            })

        return {
            "success": True,
            "records": records,
            "timeline": timeline
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.post("/save")
def save_aks_summary_data(payload: AKSSummarySavePayload):
    try:
        conn = get_db_connection()
        cur = conn.cursor()
        for item in payload.updates:
            if item.field not in ALLOWED_FIELDS:
                continue

            val = str(item.value).strip() if item.value is not None else ""
            query = f"""
                INSERT INTO equipment_aks_details (equipment_id, {item.field}, updated_at)
                VALUES (%s, %s, CURRENT_TIMESTAMP)
                ON CONFLICT (equipment_id)
                DO UPDATE SET {item.field} = EXCLUDED.{item.field}, updated_at = CURRENT_TIMESTAMP;
            """
            cur.execute(query, (item.equipment_id, val))
        conn.commit()
        cur.close()
        conn.close()
        return {"success": True, "message": f"{len(payload.updates)} updates saved successfully"}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

@router.get("/export-excel")
def export_aks_summary_excel():
    try:
        data_res = get_aks_summary_data()
        records = data_res["records"]
        timeline = data_res.get("timeline", [])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "AKS Summary"

        head_fill_main = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        head_fill_month = PatternFill(start_color="C6D9F1", end_color="C6D9F1", fill_type="solid")
        head_font = Font(name="Segoe UI", size=11, bold=True, color="000000")
        regular_font = Font(name="Segoe UI", size=10)
        center_align = Alignment(horizontal="center", vertical="center")
        left_align = Alignment(horizontal="left", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='C0C0C0'),
            right=Side(style='thin', color='C0C0C0'),
            top=Side(style='thin', color='C0C0C0'),
            bottom=Side(style='thin', color='C0C0C0')
        )

        base_headers = [
            "SN", "Asset Type", "Model", "Make", "Owner Name", "Plate No / Asset Code",
            "Date Of Purchase", "Purchase Cost", "MTC During Purchase",
            "Insurance Cost", "Reached at Site", "Which Site"
        ]

        ws.row_dimensions[1].height = 26
        for c_idx, h_text in enumerate(base_headers, 1):
            c = ws.cell(row=1, column=c_idx, value=h_text)
            c.fill = head_fill_main
            c.font = head_font
            c.alignment = center_align
            c.border = thin_border

        for t_idx, tm in enumerate(timeline):
            c = ws.cell(row=1, column=len(base_headers) + 1 + t_idx, value=tm["label"])
            c.fill = head_fill_month
            c.font = head_font
            c.alignment = center_align
            c.border = thin_border

        total_cols = len(base_headers) + len(timeline)
        curr_row = 2

        for idx, item in enumerate(records, 1):
            ws.row_dimensions[curr_row].height = 22
            p_date_val = item.get("purchase_date")
            p_date_obj = None
            if p_date_val:
                try:
                    p_date_obj = datetime.strptime(str(p_date_val)[:10], "%Y-%m-%d").date()
                except Exception:
                    p_date_obj = None

            r_date_val = item.get("reached_site_date")
            r_date_obj = None
            if r_date_val:
                try:
                    r_date_obj = datetime.strptime(str(r_date_val)[:10], "%Y-%m-%d").date()
                except Exception:
                    r_date_obj = None

            ws.cell(row=curr_row, column=1, value=idx).alignment = center_align
            ws.cell(row=curr_row, column=2, value=item.get("asset_type") or "").alignment = center_align
            ws.cell(row=curr_row, column=3, value=item.get("model") or "").alignment = center_align
            ws.cell(row=curr_row, column=4, value=item.get("make") or "").alignment = center_align
            ws.cell(row=curr_row, column=5, value=item.get("owner_name") or "").alignment = center_align
            ws.cell(row=curr_row, column=6, value=item.get("plate_no") or "").alignment = center_align

            c_pdate = ws.cell(row=curr_row, column=7, value=p_date_obj)
            c_pdate.alignment = center_align
            if p_date_obj:
                c_pdate.number_format = 'dd-mmm-yy'

            c_pcost = ws.cell(row=curr_row, column=8, value=float(item.get("purchase_cost") or 0))
            c_pcost.alignment = center_align
            c_pcost.number_format = '#,##0.00'

            c_pmtc = ws.cell(row=curr_row, column=9, value=float(item.get("mtc_during_purchase") or 0))
            c_pmtc.alignment = center_align
            c_pmtc.number_format = '#,##0.00'

            c_pins = ws.cell(row=curr_row, column=10, value=float(item.get("total_insurance") or 0))
            c_pins.alignment = center_align
            c_pins.number_format = '#,##0.00'

            c_rdate = ws.cell(row=curr_row, column=11, value=r_date_obj)
            c_rdate.alignment = center_align
            if r_date_obj:
                c_rdate.number_format = 'dd-mmm-yy'

            ws.cell(row=curr_row, column=12, value=item.get("first_site") or "").alignment = left_align

            m_mtc_list = item.get("monthly_mtc", [])
            for m_idx, m_val in enumerate(m_mtc_list):
                val_num = float(m_val or 0)
                c_m = ws.cell(row=curr_row, column=13 + m_idx, value=val_num if val_num > 0 else "")
                c_m.alignment = center_align
                if val_num > 0:
                    c_m.number_format = '#,##0.00'

            for c_num in range(1, total_cols + 1):
                cell = ws.cell(row=curr_row, column=c_num)
                cell.border = thin_border
                cell.font = regular_font

            curr_row += 1

        last_col_letter = get_column_letter(total_cols)
        ws.auto_filter.ref = f"A1:{last_col_letter}{curr_row - 1}"
        ws.sheet_view.zoomScale = 85
        ws.freeze_panes = 'F2'

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = val.strftime('%d-%b-%y') if isinstance(val, (datetime, date)) else str(val)
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"Equipment_AKS_Summary_{date.today().strftime('%Y_%m_%d')}.xlsx"
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))