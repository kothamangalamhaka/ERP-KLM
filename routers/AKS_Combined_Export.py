import os
import io
import calendar
from datetime import datetime, date
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response
import psycopg2
from psycopg2.extras import RealDictCursor
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from routers.AKS import get_aks_data
from routers.Assets_own_EQ_parts import get_assets_data
from routers.AKS_Summary import get_aks_summary_data

router = APIRouter(prefix="/py/aks-combined", tags=["AKS Combined Export"])

@router.get("/export")
def export_combined_aks_excel(scope: str = "all"):
    try:
        cur_year = date.today().year

        wb = openpyxl.Workbook()
        ws_eq = wb.active
        ws_eq.title = "EQ Assets"
        ws_assets = wb.create_sheet(title="Assets Tracker")
        ws_sum = wb.create_sheet(title="EQ Assets Exp - Summary")

        head_fill_main = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        head_fill_month = PatternFill(start_color="C6D9F1", end_color="C6D9F1", fill_type="solid")
        head_fill_sub = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
        otc_data_fill = PatternFill(start_color="FEFCE8", end_color="FEFCE8", fill_type="solid")
        rev_data_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        head_font = Font(name="Segoe UI", size=11, bold=True, color="000000")
        regular_font = Font(name="Segoe UI", size=10)
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin', color='C0C0C0'),
            right=Side(style='thin', color='C0C0C0'),
            top=Side(style='thin', color='C0C0C0'),
            bottom=Side(style='thin', color='C0C0C0')
        )

        # -------------------------------------------------------------
        # TAB 1: EQ Assets (From AKS.py)
        # -------------------------------------------------------------
        aks_res = get_aks_data()
        aks_records = aks_res.get("records", [])
        aks_timeline = aks_res.get("timeline", [])
        if scope == "current":
            tab1_timeline = [tm for tm in aks_timeline if tm["year"] == cur_year]
        else:
            tab1_timeline = aks_timeline

        ws_eq.row_dimensions[1].height = 25
        ws_eq.row_dimensions[2].height = 25
        single_headers_1 = [
            (1, "SN"), (2, "Asset Type"), (3, "Model"), (4, "Make"),
            (5, "Owner Name"), (6, "Plate No / Asset code"), (7, "Date Of Purchase"),
            (8, "Purchase Cost"), (11, "Off Site MTC")
        ]
        for c_idx, h_text in single_headers_1:
            ws_eq.merge_cells(start_row=1, start_column=c_idx, end_row=2, end_column=c_idx)
            c = ws_eq.cell(row=1, column=c_idx, value=h_text)
            c.fill = head_fill_main; c.font = head_font; c.alignment = center_align

        ws_eq.merge_cells(start_row=1, start_column=9, end_row=1, end_column=10)
        c_ins = ws_eq.cell(row=1, column=9, value="Insurance Cost")
        c_ins.fill = head_fill_main; c_ins.font = head_font; c_ins.alignment = center_align
        for sub_c, sub_t in [(9, "Off Site"), (10, "On Site")]:
            c_sub = ws_eq.cell(row=2, column=sub_c, value=sub_t)
            c_sub.fill = head_fill_sub; c_sub.font = head_font; c_sub.alignment = center_align

        col_cursor = 12
        for tm in tab1_timeline:
            ws_eq.merge_cells(start_row=1, start_column=col_cursor, end_row=1, end_column=col_cursor + 3)
            c_top = ws_eq.cell(row=1, column=col_cursor, value=tm["label"])
            c_top.fill = head_fill_month; c_top.font = head_font; c_top.alignment = center_align
            for sub_idx, sub_label in enumerate(["Site", "MTC", "OTC", "TOR"]):
                c_sub = ws_eq.cell(row=2, column=col_cursor + sub_idx, value=sub_label)
                c_sub.fill = head_fill_sub; c_sub.font = head_font; c_sub.alignment = center_align
            col_cursor += 4

        total_cols_1 = col_cursor - 1
        for r in [1, 2]:
            for c in range(1, total_cols_1 + 1):
                ws_eq.cell(row=r, column=c).border = thin_border

        curr_row = 3
        for idx, item in enumerate(aks_records, 1):
            ws_eq.row_dimensions[curr_row].height = 24
            p_date = item.get("purchase_date")
            p_date_obj = None
            if p_date:
                try: p_date_obj = datetime.strptime(str(p_date)[:10], "%Y-%m-%d").date()
                except Exception: p_date_obj = None

            ws_eq.cell(row=curr_row, column=1, value=idx).alignment = center_align
            ws_eq.cell(row=curr_row, column=2, value=item.get("asset_type") or "").alignment = center_align
            ws_eq.cell(row=curr_row, column=3, value=item.get("model") or "").alignment = center_align
            ws_eq.cell(row=curr_row, column=4, value=item.get("make") or "").alignment = center_align
            ws_eq.cell(row=curr_row, column=5, value=item.get("owner_name") or "").alignment = center_align
            ws_eq.cell(row=curr_row, column=6, value=item.get("plate_no") or "").alignment = center_align
            c_date = ws_eq.cell(row=curr_row, column=7, value=p_date_obj)
            c_date.alignment = center_align
            if p_date_obj: c_date.number_format = 'dd-mmm-yy'

            c_pcost = ws_eq.cell(row=curr_row, column=8, value=float(item.get("purchase_cost") or 0))
            c_pcost.alignment = center_align; c_pcost.number_format = '#,##0.00'

            c_off_ins = ws_eq.cell(row=curr_row, column=9, value=float(item.get("off_site_insurance") or 0))
            c_off_ins.alignment = center_align; c_off_ins.number_format = '#,##0.00'

            c_on_ins = ws_eq.cell(row=curr_row, column=10, value=float(item.get("on_site_insurance") or 0))
            c_on_ins.alignment = center_align; c_on_ins.number_format = '#,##0.00'

            c_offmtc = ws_eq.cell(row=curr_row, column=11, value=float(item.get("off_site_mtc") or 0))
            c_offmtc.alignment = center_align; c_offmtc.font = regular_font; c_offmtc.number_format = '#,##0.00'

            m_data_list = item.get("months_data", [])
            for t_idx, tm in enumerate(tab1_timeline):
                orig_idx = next((i for i, x in enumerate(aks_timeline) if x["year"] == tm["year"] and x["month"] == tm["month"]), -1)
                m_info = m_data_list[orig_idx] if orig_idx != -1 and orig_idx < len(m_data_list) else {"site": "", "mtc": 0, "otc": 0, "tor": 0}
                site_col = 12 + (t_idx * 4)
                ws_eq.cell(row=curr_row, column=site_col, value=m_info.get("site") or "").alignment = left_align

                mtc_val = float(m_info.get("mtc") or 0)
                c_m = ws_eq.cell(row=curr_row, column=site_col + 1, value=mtc_val if mtc_val > 0 else "")
                c_m.alignment = center_align
                if mtc_val > 0: c_m.number_format = '#,##0.00'

                otc_val = float(m_info.get("otc") or 0)
                c_o = ws_eq.cell(row=curr_row, column=site_col + 2, value=otc_val if otc_val > 0 else "")
                c_o.alignment = center_align; c_o.fill = otc_data_fill
                if otc_val > 0: c_o.number_format = '#,##0.00'

                tor_val = float(m_info.get("tor") or 0)
                c_t = ws_eq.cell(row=curr_row, column=site_col + 3, value=tor_val if tor_val > 0 else "")
                c_t.alignment = center_align; c_t.fill = rev_data_fill
                if tor_val > 0: c_t.number_format = '#,##0.00'

            for c_num in range(1, total_cols_1 + 1):
                cell = ws_eq.cell(row=curr_row, column=c_num)
                cell.border = thin_border
                if cell.font == Font(): cell.font = regular_font
            curr_row += 1

        ws_eq.auto_filter.ref = f"A2:{get_column_letter(total_cols_1)}{curr_row - 1}"
        ws_eq.sheet_view.zoomScale = 85
        ws_eq.freeze_panes = 'A3'

        for col in ws_eq.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = val.strftime('%d-%b-%y') if isinstance(val, (datetime, date)) else str(val)
                    max_len = max(max_len, len(val_str))
            ws_eq.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # -------------------------------------------------------------
        # TAB 2: Assets Tracker (From Assets_own_EQ_parts.py)
        # -------------------------------------------------------------
        assets_res = get_assets_data()
        assets_list = assets_res.get("assets", [])
        assets_timeline = assets_res.get("timeline", [])
        if scope == "current":
            tab2_timeline = [tm for tm in assets_timeline if tm["year"] == cur_year]
        else:
            tab2_timeline = assets_timeline

        ws_assets.row_dimensions[1].height = 25
        ws_assets.row_dimensions[2].height = 25
        base_headers_2 = [
            "SN", "Asset Type", "Asset Code", "Date of Purchase", "Cost of Purchase", "Remaining Purchase Price"
        ]
        for c_idx, h_text in enumerate(base_headers_2, 1):
            ws_assets.merge_cells(start_row=1, start_column=c_idx, end_row=2, end_column=c_idx)
            c = ws_assets.cell(row=1, column=c_idx, value=h_text)
            c.fill = head_fill_main; c.font = head_font; c.alignment = center_align

        col_cursor = len(base_headers_2) + 1
        for tm in tab2_timeline:
            ws_assets.merge_cells(start_row=1, start_column=col_cursor, end_row=1, end_column=col_cursor + 4)
            c_top = ws_assets.cell(row=1, column=col_cursor, value=tm["label"])
            c_top.fill = head_fill_month; c_top.font = head_font; c_top.alignment = center_align
            for sub_idx, sub_label in enumerate(["Site", "Plate No", "Off Site Exp", "On site", "Revenue"]):
                c_sub = ws_assets.cell(row=2, column=col_cursor + sub_idx, value=sub_label)
                c_sub.fill = head_fill_sub; c_sub.font = head_font; c_sub.alignment = center_align
            col_cursor += 5

        total_cols_2 = col_cursor - 1
        for r in [1, 2]:
            for c in range(1, total_cols_2 + 1):
                ws_assets.cell(row=r, column=c).border = thin_border

        curr_row = 3
        for idx, item in enumerate(assets_list, 1):
            ws_assets.row_dimensions[curr_row].height = 24
            p_date = item.get("purchase_date")
            p_date_obj = None
            if p_date:
                try: p_date_obj = datetime.strptime(str(p_date)[:10], "%Y-%m-%d").date()
                except Exception: p_date_obj = None

            ws_assets.cell(row=curr_row, column=1, value=idx).alignment = center_align
            ws_assets.cell(row=curr_row, column=2, value=item.get("type_name") or "").alignment = center_align
            ws_assets.cell(row=curr_row, column=3, value=item.get("asset_code") or "").alignment = center_align
            c_date = ws_assets.cell(row=curr_row, column=4, value=p_date_obj)
            c_date.alignment = center_align
            if p_date_obj: c_date.number_format = 'dd-mmm-yy'

            c_cost = ws_assets.cell(row=curr_row, column=5, value=float(item.get("purchase_cost") or 0))
            c_cost.alignment = center_align; c_cost.number_format = '#,##0.00'

            c_rem = ws_assets.cell(row=curr_row, column=6, value=float(item.get("remaining_purchase_price") or 0))
            c_rem.alignment = center_align; c_rem.number_format = '#,##0.00'

            m_data_list = item.get("months_data", [])
            for t_idx, tm in enumerate(tab2_timeline):
                orig_idx = next((i for i, x in enumerate(assets_timeline) if x["year"] == tm["year"] and x["month"] == tm["month"]), -1)
                m_info = m_data_list[orig_idx] if orig_idx != -1 and orig_idx < len(m_data_list) else {
                    "site_name": "", "plate_no": "", "off_site_exp": 0, "on_site_exp": 0, "revenue": 0
                }
                s_col = 7 + (t_idx * 5)
                ws_assets.cell(row=curr_row, column=s_col, value=m_info.get("site_name") or "").alignment = left_align
                ws_assets.cell(row=curr_row, column=s_col + 1, value=m_info.get("plate_no") or "").alignment = center_align

                off_val = float(m_info.get("off_site_exp") or 0)
                c_off = ws_assets.cell(row=curr_row, column=s_col + 2, value=off_val if off_val > 0 else "")
                c_off.alignment = center_align
                if off_val > 0: c_off.number_format = '#,##0.00'

                on_val = float(m_info.get("on_site_exp") or 0)
                c_on = ws_assets.cell(row=curr_row, column=s_col + 3, value=on_val if on_val > 0 else "")
                c_on.alignment = center_align
                if on_val > 0: c_on.number_format = '#,##0.00'

                rev_val = float(m_info.get("revenue") or 0)
                c_rev = ws_assets.cell(row=curr_row, column=s_col + 4, value=rev_val if rev_val > 0 else "")
                c_rev.alignment = center_align; c_rev.fill = rev_data_fill
                if rev_val > 0: c_rev.number_format = '#,##0.00'

            for c_num in range(1, total_cols_2 + 1):
                cell = ws_assets.cell(row=curr_row, column=c_num)
                cell.border = thin_border
                if cell.font == Font(): cell.font = regular_font
            curr_row += 1

        ws_assets.auto_filter.ref = f"A2:{get_column_letter(total_cols_2)}{curr_row - 1}"
        ws_assets.sheet_view.zoomScale = 80
        ws_assets.freeze_panes = 'A3'

        for col in ws_assets.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = val.strftime('%d-%b-%y') if isinstance(val, (datetime, date)) else str(val)
                    max_len = max(max_len, len(val_str))
            ws_assets.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # -------------------------------------------------------------
        # TAB 3: EQ Assets Exp - Summary (From AKS_Summary.py)
        # -------------------------------------------------------------
        sum_res = get_aks_summary_data()
        sum_records = sum_res.get("records", [])
        sum_timeline = sum_res.get("timeline", [])
        if scope == "current":
            tab3_timeline = [tm for tm in sum_timeline if tm["year"] == cur_year]
        else:
            tab3_timeline = sum_timeline

        base_headers_3 = [
            "SN", "Asset Type", "Model", "Make", "Plate No / Asset Code",
            "Date Of Purchase", "Purchase Cost", "MTC During Purchase",
            "Insurance Cost", "Total Cost", "Reached at Site", "Which Site"
        ]
        ws_sum.row_dimensions[1].height = 26
        for c_idx, h_text in enumerate(base_headers_3, 1):
            c = ws_sum.cell(row=1, column=c_idx, value=h_text)
            c.fill = head_fill_main; c.font = head_font; c.alignment = center_align; c.border = thin_border

        for t_idx, tm in enumerate(tab3_timeline):
            c = ws_sum.cell(row=1, column=len(base_headers_3) + 1 + t_idx, value=tm["label"])
            c.fill = head_fill_month; c.font = head_font; c.alignment = center_align; c.border = thin_border

        total_cols_3 = len(base_headers_3) + len(tab3_timeline)
        curr_row = 2
        for idx, item in enumerate(sum_records, 1):
            ws_sum.row_dimensions[curr_row].height = 22
            p_date_obj = None
            if item.get("purchase_date"):
                try: p_date_obj = datetime.strptime(str(item["purchase_date"])[:10], "%Y-%m-%d").date()
                except Exception: p_date_obj = None

            r_date_obj = None
            if item.get("reached_site_date"):
                try: r_date_obj = datetime.strptime(str(item["reached_site_date"])[:10], "%Y-%m-%d").date()
                except Exception: r_date_obj = None

            p_cost = float(item.get("purchase_cost") or 0)
            p_mtc = float(item.get("mtc_during_purchase") or 0)
            p_ins = float(item.get("total_insurance") or 0)
            tot_cost = p_cost + p_mtc + p_ins

            ws_sum.cell(row=curr_row, column=1, value=idx).alignment = center_align
            ws_sum.cell(row=curr_row, column=2, value=item.get("asset_type") or "").alignment = center_align
            ws_sum.cell(row=curr_row, column=3, value=item.get("model") or "").alignment = center_align
            ws_sum.cell(row=curr_row, column=4, value=item.get("make") or "").alignment = center_align
            ws_sum.cell(row=curr_row, column=5, value=item.get("plate_no") or "").alignment = center_align

            c_pdate = ws_sum.cell(row=curr_row, column=6, value=p_date_obj)
            c_pdate.alignment = center_align
            if p_date_obj: c_pdate.number_format = 'dd-mmm-yy'

            c_pcost = ws_sum.cell(row=curr_row, column=7, value=p_cost)
            c_pcost.alignment = center_align; c_pcost.number_format = '#,##0.00'

            c_pmtc = ws_sum.cell(row=curr_row, column=8, value=p_mtc)
            c_pmtc.alignment = center_align; c_pmtc.number_format = '#,##0.00'

            c_pins = ws_sum.cell(row=curr_row, column=9, value=p_ins)
            c_pins.alignment = center_align; c_pins.number_format = '#,##0.00'

            c_tot = ws_sum.cell(row=curr_row, column=10, value=tot_cost)
            c_tot.alignment = center_align; c_tot.number_format = '#,##0.00'; c_tot.fill = rev_data_fill

            c_rdate = ws_sum.cell(row=curr_row, column=11, value=r_date_obj)
            c_rdate.alignment = center_align
            if r_date_obj: c_rdate.number_format = 'dd-mmm-yy'

            ws_sum.cell(row=curr_row, column=12, value=item.get("first_site") or "").alignment = left_align

            m_mtc_list = item.get("monthly_mtc", [])
            for t_idx, tm in enumerate(tab3_timeline):
                orig_idx = next((i for i, x in enumerate(sum_timeline) if x["year"] == tm["year"] and x["month"] == tm["month"]), -1)
                val_num = float(m_mtc_list[orig_idx] or 0) if orig_idx != -1 and orig_idx < len(m_mtc_list) else 0.0
                c_m = ws_sum.cell(row=curr_row, column=13 + t_idx, value=val_num if val_num > 0 else "")
                c_m.alignment = center_align
                if val_num > 0: c_m.number_format = '#,##0.00'

            for c_num in range(1, total_cols_3 + 1):
                cell = ws_sum.cell(row=curr_row, column=c_num)
                cell.border = thin_border
                cell.font = regular_font
            curr_row += 1

        ws_sum.auto_filter.ref = f"A1:{get_column_letter(total_cols_3)}{curr_row - 1}"
        ws_sum.sheet_view.zoomScale = 85
        ws_sum.freeze_panes = 'F2'

        for col in ws_sum.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = cell.value
                if val:
                    val_str = val.strftime('%d-%b-%y') if isinstance(val, (datetime, date)) else str(val)
                    max_len = max(max_len, len(val_str))
            ws_sum.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"AKS_Master_Export_{scope}_{date.today().strftime('%Y_%m_%d')}.xlsx"
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))